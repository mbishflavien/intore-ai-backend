"""Text extraction utilities for various file formats"""
import base64
import io
from typing import Optional

import pdfplumber
from docx import Document


def extract_text_from_base64(base64_str: str, mime_type: str) -> str:
    """Extract text from base64 encoded file"""
    if ',' in base64_str:
        base64_str = base64_str.split(',')[1]
    
    file_bytes = base64.b64decode(base64_str)
    return extract_text_from_bytes(file_bytes, mime_type)


def extract_text_from_bytes(file_bytes: bytes, mime_type: str) -> str:
    """Extract text from file bytes"""
    
    if mime_type == "application/pdf":
        return extract_from_pdf(file_bytes)
    elif mime_type in [
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
        "application/vnd.ms-word"
    ]:
        return extract_from_docx(file_bytes)
    elif mime_type == "text/plain":
        return file_bytes.decode('utf-8')
    else:
        # Try PDF first, then fallback
        try:
            return extract_from_pdf(file_bytes)
        except:
            return extract_from_docx(file_bytes)


def extract_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF using pdfplumber"""
    text_parts = []
    
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
    
    return "\n".join(text_parts)


def extract_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX using python-docx"""
    doc = Document(io.BytesIO(file_bytes))
    text_parts = []
    
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)
    
    # Also extract tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_parts.append(cell.text)
    
    return "\n".join(text_parts)


def extract_from_file_path(file_path: str) -> str:
    """Extract text from file path (for testing)"""
    import os
    
    ext = os.path.splitext(file_path)[1].lower()
    
    with open(file_path, 'rb') as f:
        content = f.read()
    
    if ext == '.pdf':
        return extract_from_pdf(content)
    elif ext in ['.docx', '.doc']:
        return extract_from_docx(content)
    elif ext == '.txt':
        return content.decode('utf-8')
    else:
        raise ValueError(f"Unsupported file format: {ext}")
